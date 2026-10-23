import sys
from pathlib import Path

import joblib
import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.batch_predict_workbooks import (  # noqa: E402
    FEEDBACK_HEADERS, FEATURE_HEADERS, MODEL_FEATURES, load_and_validate,
    map_probabilities, report_rows, write_report,
)


def _workbooks(tmp_path, *, rows=2, duplicate=False, mismatch=False, invalid_numeric=False, missing_feedback=False):
    feedback_path = tmp_path / "feedback.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Teacher Feedback"; sheet.append(FEEDBACK_HEADERS)
    features_path = tmp_path / "features.xlsx"
    feature_book = Workbook(); feature_sheet = feature_book.active; feature_sheet.title = "Prototype Features"; feature_sheet.append(FEATURE_HEADERS)
    for i in range(rows):
        student_id = f"s{i+1}"
        feedback = [student_id, f"Student {i+1}", "Grade 1", "lesson-1", "Song Book", "High", "Medium", "High", "Good", "None identified", "Visual learner", "", "", "Submitted", "", "", "teacher"]
        if missing_feedback: feedback[8] = ""
        sheet.append(feedback)
        feature_sheet.append([student_id, f"Student {i+1}", "Grade 1", "other-lesson" if mismatch else "lesson-1", 1.5, 5, 80, 4, 2, 50])
    if duplicate:
        sheet.append(["s1", "Duplicate", "Grade 1", "lesson-1", "Song Book", "High", "Medium", "High", "Good", "None identified", "Visual learner", "", "", "Submitted", "", "", "teacher"])
    if invalid_numeric: feature_sheet["F2"] = 1.5
    book.save(feedback_path); feature_book.save(features_path)
    return feedback_path, features_path


def test_prototype_schema_matches_exact_java_service_and_candidate_order():
    model = joblib.load(ROOT / "models/candidates/binary-10-feature-prototype/random_forest.joblib")
    assert list(model.feature_names_in_) == MODEL_FEATURES
    java = (ROOT.parent / "backend/src/main/java/lk/smartplay/adaptive/service/BinaryPrototypeFeatureService.java").read_text()
    assert all(f'"{name}"' in java for name in MODEL_FEATURES)
    assert list(model.classes_) == ["Needs Practice", "Successful"]
    assert model.named_steps["preprocessor"].named_transformers_["categorical"].handle_unknown == "ignore"


def test_probability_mapping_uses_model_class_order_and_only_supported_classes():
    assert map_probabilities(["Successful", "Needs Practice"], [0.72, 0.28]) == {"Successful": 0.72, "Needs Practice": 0.28}
    with pytest.raises(ValueError, match="exactly its two supported classes"):
        map_probabilities(["Needs Practice", "Waiting"], [0.5, 0.5])


def test_exact_student_id_join_and_consistent_lesson_worksheet_ids(tmp_path):
    feedback, features = _workbooks(tmp_path)
    joined = load_and_validate(feedback, features)
    assert [item["feedback"]["Student ID"] for item in joined] == ["s1", "s2"]
    feedback, features = _workbooks(tmp_path, mismatch=True)
    with pytest.raises(ValueError, match="Lesson ID and Worksheet ID"):
        load_and_validate(feedback, features)
    feedback, features = _workbooks(tmp_path, duplicate=True)
    with pytest.raises(ValueError, match="duplicate Student ID"):
        load_and_validate(feedback, features)
    feedback, features = _workbooks(tmp_path)
    feature_book = load_workbook(features); feature_book["Prototype Features"]["A2"] = "other-id"; feature_book.save(features)
    with pytest.raises(ValueError, match="Exact Student ID mismatch"):
        load_and_validate(feedback, features)


def test_builds_exact_ten_features_from_workbook_and_saved_feedback_categories(tmp_path):
    feedback, features = _workbooks(tmp_path, rows=1)
    item = load_and_validate(feedback, features)[0]
    assert list(item["features"]) == MODEL_FEATURES
    assert item["features"]["Average_Attempts"] == 1.5
    assert item["features"]["Participation"] == "High"
    assert item["features"]["Attention_Level"] == "Medium"
    assert item["features"]["Collaboration"] == "Good"
    assert item["features"]["Learning_Difficulty"] == "None identified"
    assert "Confidence" not in item["features"] and "Learning_Style" not in item["features"]
    assert item["issues"] == []


def test_missing_or_invalid_numeric_and_feedback_fields_are_identified(tmp_path):
    feedback, features = _workbooks(tmp_path, rows=1, invalid_numeric=True, missing_feedback=True)
    item = load_and_validate(feedback, features)[0]
    assert "Total_Activities (must be a whole number)" in item["issues"]
    assert "Collaboration (missing or invalid prototype category)" in item["issues"]


def test_batch_report_keeps_prediction_fields_blank_when_saved_snapshot_is_unconfirmed(tmp_path):
    feedback, features = _workbooks(tmp_path, rows=2)
    joined = load_and_validate(feedback, features)
    rows = report_rows(joined, "prototype-hash", "Authenticated snapshot verification required")
    assert len(rows) == 2
    assert all(row["Predicted class"] == "" and row["Needs Practice probability"] == "" and row["Successful probability"] == "" for row in rows)
    assert all(row["Prediction status"] == "BLOCKED_AUTHENTICATION_REQUIRED" for row in rows)
    output = tmp_path / "prediction-report.xlsx"
    model = joblib.load(ROOT / "models/candidates/binary-10-feature-prototype/random_forest.joblib")
    write_report(rows, output, "prototype-hash", list(model.classes_), MODEL_FEATURES, [], "service health verified")
    saved = load_workbook(output, data_only=True)
    sheet = saved["Batch Predictions"]
    assert sheet.max_row == 3
    assert [cell.value for cell in sheet[1]] == ["Student ID", "Student Name", "Grade", "Lesson ID", "Worksheet ID", "Model name and version", "Predicted class", "Needs Practice probability", "Successful probability", "Prediction timestamp", "Prediction status", "Missing or invalid input fields"]
    assert sheet["K2"].value == "BLOCKED_AUTHENTICATION_REQUIRED"
    assert saved["Model Notes"]["B4"].value == "Needs Practice, Successful"
