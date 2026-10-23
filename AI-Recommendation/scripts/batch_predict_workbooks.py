"""Preflight prototype workbook inputs and write a report without bypassing saved state.

Predictions are intentionally gated on confirming snapshots through the authenticated
backend workflow. This command does not write to MongoDB or call the model directly.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
from typing import Any

import joblib
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
MODEL_PATH = ROOT / "models/candidates/binary-10-feature-prototype/random_forest.joblib"
MODEL_FEATURES = ["Average_Attempts", "Total_Activities", "Completion_Rate", "Participation", "Attention_Level", "Collaboration", "Learning_Difficulty", "Historical_Lessons", "Completed_Historical_Lessons", "Historical_Completion_Rate"]
NUMERIC_FEATURES = ["Average_Attempts", "Total_Activities", "Completion_Rate", "Historical_Lessons", "Completed_Historical_Lessons", "Historical_Completion_Rate"]
FEEDBACK_COLUMNS = {"Participation": "Participation", "Attention_Level": "Attention Level", "Collaboration": "Collaboration", "Learning_Difficulty": "Learning Difficulty"}
FEEDBACK_HEADERS = ["Student ID", "Student Name", "Grade", "Lesson ID", "Lesson Name", "Participation", "Attention Level", "Confidence", "Collaboration", "Learning Difficulty", "Learning Style", "Improvement Observed", "Feedback / Next Step", "Feedback Status", "Feedback Submission Date and Time", "Last Updated Date and Time", "Teacher ID"]
FEATURE_HEADERS = ["Student ID", "Student Name", "Grade", "Worksheet ID", *NUMERIC_FEATURES]
REPORT_HEADERS = ["Student ID", "Student Name", "Grade", "Lesson ID", "Worksheet ID", "Model name and version", "Predicted class", "Needs Practice probability", "Successful probability", "Prediction timestamp", "Prediction status", "Missing or invalid input fields"]


def _load_rows(path: Path, sheet_name: str, headers: list[str]) -> list[dict[str, Any]]:
    book = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in book.sheetnames:
        raise ValueError(f"Missing worksheet: {sheet_name}")
    sheet = book[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    if list(next(iterator, ())) != headers:
        raise ValueError(f"{path.name}: unexpected headers or feature order")
    rows = []
    for number, cells in enumerate(iterator, 2):
        if all(value is None for value in cells):
            continue
        row = dict(zip(headers, cells))
        student_id = row.get("Student ID")
        if not isinstance(student_id, str) or not student_id or student_id.strip() != student_id:
            raise ValueError(f"{path.name}: invalid Student ID at row {number}")
        rows.append(row)
    book.close()
    ids = [row["Student ID"] for row in rows]
    if len(ids) != len(set(ids)):
        raise ValueError(f"{path.name}: duplicate Student ID")
    return rows


def _accepted_categories(model) -> dict[str, set[str]]:
    pipeline = model.named_steps["preprocessor"]
    columns = next(columns for name, _, columns in pipeline.transformers_ if name == "categorical")
    encoder = pipeline.named_transformers_["categorical"]
    return {name: set(values.tolist()) for name, values in zip(columns, encoder.categories_)}


def map_probabilities(class_labels: list[str], probability_values: list[float]) -> dict[str, float]:
    """Associate probabilities only with the actual class order returned by the model."""
    if len(class_labels) != len(probability_values) or set(class_labels) != {"Needs Practice", "Successful"}:
        raise ValueError("Prototype must return probabilities for exactly its two supported classes")
    if any(not math.isfinite(value) or value < 0 or value > 1 for value in probability_values):
        raise ValueError("Prototype returned an invalid probability")
    mapped = dict(zip(class_labels, probability_values))
    if abs(sum(mapped.values()) - 1.0) > 1e-9:
        raise ValueError("Prototype probabilities must sum to one")
    return mapped


def load_and_validate(feedback_path: Path, features_path: Path, model=None) -> list[dict[str, Any]]:
    feedback = _load_rows(feedback_path, "Teacher Feedback", FEEDBACK_HEADERS)
    feature_rows = _load_rows(features_path, "Prototype Features", FEATURE_HEADERS)
    by_id = {row["Student ID"]: row for row in feature_rows}
    feedback_ids = {row["Student ID"] for row in feedback}
    if feedback_ids != set(by_id):
        raise ValueError(f"Exact Student ID mismatch: no feature row for {sorted(feedback_ids-set(by_id))}; no feedback row for {sorted(set(by_id)-feedback_ids)}")
    lessons = {row["Lesson ID"] for row in feedback}
    worksheets = {row["Worksheet ID"] for row in feature_rows}
    if len(lessons) != 1 or len(worksheets) != 1 or lessons != worksheets:
        raise ValueError(f"Lesson ID and Worksheet ID must match consistently: lessons={sorted(map(str, lessons))}, worksheets={sorted(map(str, worksheets))}")
    model = model or joblib.load(MODEL_PATH)
    categories = _accepted_categories(model)
    joined = []
    for fb in feedback:
        numeric = by_id[fb["Student ID"]]
        issues = []
        if fb["Student Name"] != numeric["Student Name"]: issues.append("Student Name (workbooks disagree)")
        if fb["Grade"] != numeric["Grade"]: issues.append("Grade (workbooks disagree)")
        for name in NUMERIC_FEATURES:
            value = numeric[name]
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
                issues.append(f"{name} (must be finite and non-negative)")
        for name in ["Total_Activities", "Historical_Lessons", "Completed_Historical_Lessons"]:
            value = numeric[name]
            if isinstance(value, (int, float)) and math.isfinite(value) and value != int(value):
                issues.append(f"{name} (must be a whole number)")
        for name in ["Completion_Rate", "Historical_Completion_Rate"]:
            value = numeric[name]
            if isinstance(value, (int, float)) and math.isfinite(value) and value > 100:
                issues.append(f"{name} (must be 0–100)")
        if all(isinstance(numeric[name], (int, float)) and math.isfinite(numeric[name]) for name in ["Historical_Lessons", "Completed_Historical_Lessons"]) and numeric["Completed_Historical_Lessons"] > numeric["Historical_Lessons"]:
            issues.append("Completed_Historical_Lessons (exceeds Historical_Lessons)")
        if fb["Feedback Status"] != "Submitted":
            issues.append("Feedback Status (saved Submitted feedback required)")
        for feature, column in FEEDBACK_COLUMNS.items():
            value = fb[column]
            if not isinstance(value, str) or value not in categories[feature]:
                issues.append(f"{column} (missing or invalid prototype category)")
        features = {name: numeric[name] for name in NUMERIC_FEATURES}
        for feature, column in FEEDBACK_COLUMNS.items():
            features[feature] = fb[column]
        ordered = {name: features[name] for name in MODEL_FEATURES}
        joined.append({"feedback": fb, "numeric": numeric, "features": ordered, "issues": issues})
    return joined


def report_rows(joined: list[dict[str, Any]], model_version: str, blocked_reason: str) -> list[dict[str, Any]]:
    results = []
    for item in joined:
        fb, numeric = item["feedback"], item["numeric"]
        missing = list(item["issues"])
        missing.append(blocked_reason)
        results.append({
            "Student ID": fb["Student ID"], "Student Name": fb["Student Name"], "Grade": fb["Grade"],
            "Lesson ID": fb["Lesson ID"], "Worksheet ID": numeric["Worksheet ID"],
            "Model name and version": f"10-feature Random Forest prototype · SHA-256 {model_version}",
            "Predicted class": "", "Needs Practice probability": "", "Successful probability": "", "Prediction timestamp": "",
            "Prediction status": "BLOCKED_AUTHENTICATION_REQUIRED", "Missing or invalid input fields": "; ".join(missing),
        })
    return results


def write_report(rows: list[dict[str, Any]], destination: Path, model_version: str, classes: list[str], feature_order: list[str], limitations: list[str], health_note: str) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Batch Predictions"
    sheet.append(REPORT_HEADERS)
    for row in rows:
        sheet.append([row.get(name, "") for name in REPORT_HEADERS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(REPORT_HEADERS))}{max(1, len(rows)+1)}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col, width in enumerate([27, 28, 14, 34, 34, 82, 24, 28, 28, 28, 42, 90], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    notes = book.create_sheet("Model Notes")
    notes.append(("Model", "10-feature Random Forest prototype"))
    notes.append(("Artifact", str(MODEL_PATH.relative_to(ROOT))))
    notes.append(("Model version SHA-256", model_version))
    notes.append(("Actual model classes", ", ".join(classes)))
    notes.append(("Exact feature order", ", ".join(feature_order)))
    notes.append(("Service health", health_note))
    notes.append(("Workbook rows validated", len(rows)))
    notes.append(("Successful predictions", sum(row["Prediction status"] == "PREDICTED" for row in rows)))
    notes.append(("Blocked students", sum(row["Prediction status"].startswith("BLOCKED") for row in rows)))
    notes.append(("Confirmed missing database records", "Unknown; authenticated student/feedback/snapshot reads were unavailable."))
    notes.append(("Inference failures", 0))
    notes.append(("Inference calls attempted", 0))
    notes.append(("Inference status", "No prediction was produced because /api/auth/me returned 204; the protected import, feedback verification, and prototype prediction workflows require an authenticated teacher session."))
    notes.append(("Candidate limitation", "Exploratory prototype; not validated for student decisions. Results, if later run, remain experimental and are not verified student outcomes."))
    for limitation in limitations:
        notes.append(("Validation limitation", limitation))
    notes.column_dimensions["A"].width = 34
    notes.column_dimensions["B"].width = 120
    for row in notes.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--feedback", type=Path, required=True)
    parser.add_argument("--features", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--blocked-reason", default="Authenticated teacher session required to verify/import the saved snapshot and confirm lesson-scoped teacher feedback; protected prediction call not made.")
    parser.add_argument("--health-note", default="MongoDB ping succeeded; Spring CSRF endpoint returned HTTP 200; AI /health returned status ok and binary_prototype_available=true. /api/auth/me returned HTTP 204 (no authenticated teacher session).")
    args = parser.parse_args()
    model = joblib.load(MODEL_PATH)
    model_version = hashlib.sha256(MODEL_PATH.read_bytes()).hexdigest()
    joined = load_and_validate(args.feedback, args.features, model)
    rows = report_rows(joined, model_version, args.blocked_reason)
    report_path = MODEL_PATH.parent / "prototype_report.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    write_report(rows, args.output, model_version, [str(value) for value in model.classes_], list(model.feature_names_in_), report.get("limitations", []), args.health_note)
    print(json.dumps({"studentsValidated": len(joined), "predicted": 0, "blocked": len(rows), "output": str(args.output)}, indent=2))


if __name__ == "__main__":
    main()
